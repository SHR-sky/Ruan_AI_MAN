"""
将 dataset/ 下的原始数据转化为知识库 JSON + CSV 格式
输出到 backend/data/knowledge_base/
"""
import os
import json
import csv
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from docx import Document


DATASET_DIR = os.path.join(os.path.dirname(__file__), "..", "dataset")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "backend", "data", "knowledge_base")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def find_files(ext):
    results = []
    for root, dirs, files in os.walk(DATASET_DIR):
        for f in files:
            if f.endswith(ext):
                results.append(os.path.join(root, f))
    return results


def extract_xlsx_knowledge():
    """从游客数据 xlsx 中提取唯一景点描述"""
    xlsx_files = find_files(".xlsx")
    if not xlsx_files:
        print("No xlsx file found!")
        return [], []

    xlsx_path = xlsx_files[0]
    print(f"Reading: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active

    attractions = {}
    faq_candidates = []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        name = str(row[4]).strip() if row[4] else ""
        content = str(row[5]).strip() if row[5] else ""
        atype = str(row[6]).strip() if row[6] else ""

        if not name or not content:
            continue

        # Keep the longest/most complete description per attraction
        if name not in attractions or len(content) > len(attractions[name]["content"]):
            attractions[name] = {
                "name": name,
                "type": atype,
                "content": content,
            }

        # Collect some FAQ-style records from visitor questions
        if i < 1000 and len(content) > 20:
            faq_candidates.append({
                "question": f"介绍一下{name}",
                "answer": content[:500],
                "category": atype,
            })

        if (i + 1) % 50000 == 0:
            print(f"  Processed {i+1} rows...")

    wb.close()
    print(f"  Done: {len(attractions)} unique attractions extracted")

    knowledge_items = []
    for a in attractions.values():
        item = {
            "id": f"attraction_{len(knowledge_items)+1:04d}",
            "name": a["name"],
            "type": a["type"],
            "content": a["content"],
            "keywords": generate_keywords(a["name"], a["type"], a["content"]),
            "source": "xlsx_tourist_data",
        }
        knowledge_items.append(item)

    return knowledge_items, faq_candidates


def extract_docx_knowledge():
    """从 docx 文件中提取历史文化知识"""
    docx_files = find_files(".docx")
    items = []

    for docx_path in docx_files:
        filename = os.path.basename(docx_path)
        print(f"Reading: {filename}")
        doc = Document(docx_path)

        current_section = ""
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if len(text) < 30 and para.style.name.startswith("Heading"):
                current_section = text
            elif len(text) > 50:
                items.append({
                    "id": f"culture_{len(items)+1:04d}",
                    "name": current_section or filename,
                    "type": "历史文化",
                    "content": text,
                    "keywords": generate_keywords("", "历史文化", text),
                    "source": filename,
                })

    print(f"  Extracted {len(items)} cultural knowledge items")
    return items


def generate_keywords(name, atype, content):
    """基于名称和类型生成关键词"""
    words = []
    if name:
        words.append(name)
    if atype:
        words.append(atype)

    # Extract location-related terms
    location_indicators = ["位于", "坐落", "地处", "上海", "江苏", "浙江", "景区"]
    for indicator in location_indicators:
        if indicator in content and indicator not in words:
            words.append(indicator)

    return ",".join(words[:10])


def build_faq_from_knowledge(knowledge_items):
    """基于知识库自动生成 FAQ"""
    faq_templates = [
        ("{name}在哪里", "{name}的相关信息请参考景区导览。"),
        ("介绍一下{name}", "{content_preview}"),
        ("{name}有什么特色", "{name}是{type}类型的景点。{content_preview}"),
        ("{name}好玩吗", "{content_preview}"),
        ("{name}是什么类型的景点", "{name}属于{type}类型。"),
    ]

    faqs = []
    for item in knowledge_items:
        name = item["name"]
        preview = item["content"][:200]
        atype = item["type"]
        for q_template, a_template in faq_templates:
            question = q_template.format(name=name)
            answer = a_template.format(name=name, type=atype, content_preview=preview)
            faqs.append({
                "id": f"faq_{len(faqs)+1:05d}",
                "question": question,
                "answer": answer,
                "category": atype,
            })

    return faqs


def save_json(filename, data):
    path = os.path.join(OUTPUT_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Saved: {path} ({len(data)} items)")
    return path


def save_csv(filename, fieldnames, rows):
    path = os.path.join(OUTPUT_DIR, filename)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})
    print(f"Saved: {path} ({len(rows)} rows)")


def main():
    print("=" * 60)
    print("Converting dataset to knowledge base...")
    print("=" * 60)

    # 1. Extract from xlsx
    xlsx_kb, faq_candidates = extract_xlsx_knowledge()

    # 2. Extract from docx
    docx_kb = extract_docx_knowledge()

    # 3. Merge all knowledge
    all_knowledge = xlsx_kb + docx_kb

    # 4. Auto-generate FAQ
    auto_faqs = build_faq_from_knowledge(xlsx_kb)
    merged_faqs = auto_faqs + faq_candidates

    # 5. Save as JSON
    save_json("knowledge_base.json", all_knowledge)
    save_json("faq.json", merged_faqs)

    # 6. Save as CSV
    kb_fieldnames = ["id", "name", "type", "content", "keywords", "source"]
    save_csv("knowledge_base.csv", kb_fieldnames, all_knowledge)

    faq_fieldnames = ["id", "question", "answer", "category"]
    save_csv("faq.csv", faq_fieldnames, merged_faqs)

    # 7. Stats
    print()
    print("=" * 60)
    print("Knowledge Base Summary:")
    print(f"  Total knowledge items: {len(all_knowledge)}")
    print(f"  Total FAQ items: {len(merged_faqs)}")
    kb_size = sum(len(k["content"]) for k in all_knowledge)
    print(f"  Total content characters: {kb_size:,}")
    print(f"  Output directory: {OUTPUT_DIR}")
    print("=" * 60)


if __name__ == "__main__":
    main()
